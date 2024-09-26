import os
import nltk
from PyPDF2 import PdfReader
from docx import Document
from nltk.tokenize import word_tokenize
from nltk.probability import FreqDist
from nltk.corpus import words
from fpdf import FPDF
from nltk.metrics import edit_distance

# Função para verificar e baixar pacotes do NLTK, se necessário
def verificar_nltk():
    try:
        nltk.data.find('corpora/words.zip')
    except LookupError:
        nltk.download('words')

    try:
        nltk.data.find('tokenizers/punkt.zip')
    except LookupError:
        nltk.download('punkt')

# Chame a função para verificar e baixar os pacotes, se necessário
verificar_nltk()

def analisar_texto(texto):
    palavras = word_tokenize(texto.lower())
    erros = {palavra: sugerir_correcao(palavra) for palavra in palavras if palavra not in words.words()}
    return erros

def sugerir_correcao(palavra):
    palavras_validas = set(words.words())
    sugeridas = [(p, edit_distance(palavra, p)) for p in palavras_validas]
    sugeridas.sort(key=lambda x: x[1])
    return [s[0] for s in sugeridas[:5]]

def importar_arquivo(caminho):
    if caminho.endswith('.pdf'):
        return ler_pdf(caminho)
    elif caminho.endswith('.docx'):
        return ler_docx(caminho)
    elif caminho.endswith('.txt'):
        return ler_txt(caminho)
    else:
        raise ValueError("Formato de arquivo não suportado.")

def ler_pdf(caminho):
    with open(caminho, 'rb') as arquivo:
        leitor = PdfReader(arquivo)
        texto = ""
        for pagina in leitor.pages:
            texto += pagina.extract_text()
    return texto

def ler_docx(caminho):
    doc = Document(caminho)
    texto = '\n'.join([paragrafo.text for paragrafo in doc.paragraphs])
    return texto

def ler_txt(caminho):
    with open(caminho, 'r', encoding='utf-8') as arquivo:
        return arquivo.read()

def gerar_relatorio(erros, caminho_saida):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    pdf.cell(200, 10, txt="Relatório de Erros Ortográficos", ln=True, align='C')
    pdf.ln(10)

    for erro, sugestoes in erros.items():
        pdf.cell(0, 10, txt=f"Erro: {erro} | Sugestões: {', '.join(sugestoes)}", ln=True)

    pdf.output(caminho_saida)

def analisar_frequencia(texto):
    palavras = word_tokenize(texto.lower())
    frequencia = FreqDist(palavras)
    return frequencia.most_common(10)

def salvar_historico(erros, frequencia):
    with open('historico.txt', 'a') as f:
        f.write(f"Erros: {erros}, Frequência: {frequencia}\n")
        
    with open('falhas.txt', 'w') as f:
        for erro, sugestoes in erros.items():
            f.write(f"Erro: {erro} | Sugestões: {', '.join(sugestoes)}\n")

if __name__ == "__main__":
    caminho = input("Insira o caminho do arquivo: ")
    texto = importar_arquivo(caminho)
    erros_encontrados = analisar_texto(texto)
    print("Erros encontrados e sugestões:", erros_encontrados)

    # Analisar frequência
    frequencia = analisar_frequencia(texto)
    print("Frequência das palavras:", frequencia)

    # Salvar histórico
    salvar_historico(erros_encontrados, frequencia)



